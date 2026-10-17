"""Wrap page images into a 16:9 PPTX, one full-bleed picture per slide.

Needs ``python-pptx`` (imported lazily). Feed it the CLEAN pages, not the watermarked PDF set.
"""
import io
from pathlib import Path


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _add_fade(slide):
    """python-pptx has no transition API: write ``<p:transition><p:fade/></p:transition>`` into the XML."""
    from lxml import etree
    sld = slide._element
    for old in sld.findall("{%s}transition" % P_NS):
        sld.remove(old)
    tr = etree.SubElement(sld, "{%s}transition" % P_NS, spd="med")
    etree.SubElement(tr, "{%s}fade" % P_NS)
    clr = sld.find("{%s}clrMapOvr" % P_NS)
    if clr is not None:  # schema order: cSld, clrMapOvr, transition, timing
        clr.addnext(tr)


def build_pptx(images, out, width_in=13.333, height_in=7.5, transition=None, jpeg_quality=None):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError("build_pptx needs python-pptx: pip install python-pptx") from exc
    if transition not in (None, "fade"):
        raise ValueError(f"transition must be 'fade' or None, got {transition!r}")
    for p in images:
        if not Path(p).exists():
            raise FileNotFoundError(p)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width_in), Inches(height_in)
    blank = prs.slide_layouts[6]
    for p in images:
        s = prs.slides.add_slide(blank)
        if jpeg_quality:
            from PIL import Image
            buf = io.BytesIO()
            with Image.open(p) as im:
                im.convert("RGB").save(buf, "JPEG", quality=int(jpeg_quality), optimize=True)
            buf.seek(0)
            src = buf
        else:
            src = str(p)
        s.shapes.add_picture(src, 0, 0, width=prs.slide_width, height=prs.slide_height)
        if transition == "fade":
            _add_fade(s)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return len(images)
